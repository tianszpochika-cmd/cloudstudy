import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_layout)
txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11.333), Inches(2))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "EnglishFlow 产品与商业化方案"
p.font.size = Pt(36)
p.font.bold = True
prs.save("test_presentation.pptx")
print("pptx created:", os.path.exists("test_presentation.pptx"))
